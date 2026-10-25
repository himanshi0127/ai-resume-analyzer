"""
Resume Parser Module
Extracts text and identifies contact information,
skills, education, and resume sections.
"""

import re
import pdfplumber
import docx


SKILL_DB = [
    "Python", "Java", "C++", "C", "JavaScript", "TypeScript",
    "SQL", "HTML", "CSS", "React", "Node.js", "Express",
    "Dart", "Django", "Flask", "MongoDB", "MySQL", "PostgreSQL",
    "Machine Learning", "Deep Learning", "TensorFlow", "PyTorch",
    "Pandas", "NumPy", "scikit-learn", "Power BI", "Excel",
    "Tableau", "AWS", "Azure", "Docker", "Kubernetes", "Git",
    "Linux", "REST API", "Statistics", "Data Analysis", "Kotlin",
    "Android Studio", "XML", "Network Security", "SIEM",
    "Cybersecurity", "Communication", "CI/CD", "Problem Solving",
    "Teamwork", "PHP", "OOP", "Data Structure"
]

EDUCATION_KEYWORDS = [
    "b.tech", "btech", "bachelor", "b.e", "m.tech", "mtech",
    "master", "mba", "bca", "mca", "phd", "b.sc", "m.sc",
    "diploma", "12th", "higher secondary"
]

# Each section uses a consistent key expected by the scorer.
SECTION_KEYWORDS = {
    "projects": [
        "projects",
        "project",
        "academic projects",
        "academic project",
        "personal projects",
        "personal project",
        "my projects",
        "college projects",
        "key projects",
        "project work"
    ],
    "experience": [
        "experience",
        "work experience",
        "professional experience",
        "work history",
        "employment history",
        "internship",
        "internships",
        "training",
        "industrial training"
    ],
    "certifications": [
        "certifications",
        "certification",
        "certificates",
        "certificate",
        "licenses and certifications",
        "courses and certifications",
        "training and certifications"
    ],
    "achievements": [
        "achievements",
        "achievement",
        "awards",
        "honors and awards",
        "accomplishments",
        "academic achievements"
    ],
    "summary": [
        "summary",
        "professional summary",
        "profile summary",
        "career objective",
        "objective",
        "about me",
        "profile",
        "professional profile",
        "personal statement"
    ]
}


def extract_text_from_pdf(file_path):
    text_parts = []

    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)

    return "\n".join(text_parts)


def extract_text_from_docx(file_path):
    document = docx.Document(file_path)

    parts = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    # Also extract text from tables.
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip() for cell in row.cells
            )
            if row_text.strip():
                parts.append(row_text)

    return "\n".join(parts)


def extract_text(file_path):
    file_path_lower = file_path.lower()

    if file_path_lower.endswith(".pdf"):
        return extract_text_from_pdf(file_path)

    if file_path_lower.endswith(".docx"):
        return extract_text_from_docx(file_path)

    raise ValueError(
        "Unsupported format. Please upload PDF or DOCX."
    )


def extract_email(text):
    match = re.search(
        r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
        text
    )

    return match.group(0) if match else None


def extract_phone(text):
    """Extract a plausible phone number and reject obvious placeholders."""
    candidates = re.findall(
        r"(?<!\d)(?:\+?\d{1,3}[-\s.]?)?"
        r"(?:\(?\d{3}\)?[-\s.]?)?\d{3}[-\s.]?\d{4}(?!\d)",
        text
    )

    for candidate in candidates:
        digits = re.sub(r"\D", "", candidate)

        if len(digits) == 12 and digits.startswith("91"):
            digits = digits[2:]

        if len(digits) != 10:
            continue

        if len(set(digits)) == 1:
            continue

        if digits in (
            "0123456789",
            "1234567890",
            "9876543210"
        ):
            continue

        return candidate.strip()

    return None


def extract_name(text):
    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    excluded_terms = [
        "resume", "curriculum vitae", "phone", "email",
        "linkedin", "github", "objective", "projects",
        "experience", "education", "skills", "certifications"
    ]

    for line in lines[:8]:
        lower = line.lower()

        if "@" in line or any(char.isdigit() for char in line):
            continue

        if any(term in lower for term in excluded_terms):
            continue

        if len(line.split()) <= 5 and len(line) <= 60:
            return line

    return "Not Found"


def extract_skills(text):
    found = []
    text_lower = text.lower()

    for skill in SKILL_DB:
        pattern = (
            r"(?<!\w)" + re.escape(skill.lower()) + r"(?!\w)"
        )

        if re.search(pattern, text_lower):
            found.append(skill)

    return sorted(set(found))


def extract_education(text):
    found = []
    text_lower = text.lower()

    for keyword in EDUCATION_KEYWORDS:
        pattern = (
            r"(?<!\w)" + re.escape(keyword.lower()) + r"(?!\w)"
        )

        if re.search(pattern, text_lower):
            found.append(keyword.upper())

    return sorted(set(found))


def normalize_heading(line):
    """Clean a possible section-heading line."""
    line = line.strip().lower()

    # Remove bullets or numbering at the beginning.
    line = re.sub(r"^\s*(?:[-*•▪◦]+|\d+[.)])\s*", "", line)

    # Normalize repeated whitespace.
    line = re.sub(r"\s+", " ", line)

    return line.strip(" \t:|—–-")


def extract_sections(text):
    """
    Detect section headings by examining individual lines.
    Avoid counting a section just because its keyword occurs
    somewhere in a normal sentence.
    """
    found = set()

    aliases = [
        (section, alias)
        for section, keywords in SECTION_KEYWORDS.items()
        for alias in keywords
    ]

    # Check longer aliases first to handle specific headings clearly.
    aliases.sort(key=lambda item: len(item[1]), reverse=True)

    for line in text.splitlines():
        cleaned = normalize_heading(line)

        if not cleaned:
            continue

        for section, alias in aliases:
            alias_pattern = re.escape(alias.lower())
            alias_pattern = alias_pattern.replace(r"\ ", r"\s+")

            # Accept a heading by itself, or a heading followed
            # by a colon, vertical bar, or dash and inline content.
            pattern = (
                r"^\s*" + alias_pattern +
                r"\s*(?:(?:[:|—–-])\s*.*)?$"
            )

            if re.match(pattern, cleaned, flags=re.IGNORECASE):
                found.add(section)
                break

    return sorted(found)


def parse_resume(file_path):
    raw_text = extract_text(file_path)

    return {
        "name": extract_name(raw_text),
        "email": extract_email(raw_text),
        "phone": extract_phone(raw_text),
        "skills": extract_skills(raw_text),
        "education": extract_education(raw_text),
        "sections": extract_sections(raw_text),
        "raw_text": raw_text
    }